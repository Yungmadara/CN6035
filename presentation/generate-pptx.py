# -*- coding: utf-8 -*-
"""
Generates CN6035_Presentation.pptx — 12 slides covering the assignment requirements.
Run: python presentation/generate-pptx.py
"""
import sys, io, os
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from copy import deepcopy

# Color palette (matches app theme)
BG_DARK = RGBColor(0x1a, 0x1a, 0x2e)       # navy
BG_CARD = RGBColor(0x16, 0x21, 0x3e)       # darker navy
ACCENT = RGBColor(0xe0, 0xc0, 0x68)        # gold
TEXT_WHITE = RGBColor(0xff, 0xff, 0xff)
TEXT_GRAY = RGBColor(0xaa, 0xaa, 0xaa)
ACCENT_GREEN = RGBColor(0x4a, 0xde, 0x80)
ACCENT_RED = RGBColor(0xff, 0x6b, 0x6b)
LINE_BLUE = RGBColor(0x0f, 0x34, 0x60)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H

blank_layout = prs.slide_layouts[6]


def set_bg(slide, color=BG_DARK):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.line.fill.background()
    bg.fill.solid()
    bg.fill.fore_color.rgb = color
    bg.shadow.inherit = False
    # Push to back
    spTree = bg._element.getparent()
    spTree.remove(bg._element)
    spTree.insert(2, bg._element)
    return bg


def add_text(slide, x, y, w, h, text, size=18, color=TEXT_WHITE, bold=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font='Calibri'):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.font.bold = bold
    r.font.name = font
    return tb


def add_paragraph(tf, text, size=14, color=TEXT_WHITE, bold=False, bullet=False, indent=0):
    p = tf.add_paragraph()
    if bullet:
        p.level = indent
        # Add bullet manually
        text = '• ' + text
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.font.bold = bold
    r.font.name = 'Calibri'
    return p


def add_box(slide, x, y, w, h, fill=BG_CARD, line=LINE_BLUE, line_w=1.5):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    box.fill.solid()
    box.fill.fore_color.rgb = fill
    box.line.color.rgb = line
    box.line.width = Pt(line_w)
    box.shadow.inherit = False
    return box


def accent_bar(slide, x, y, w=Inches(0.08), h=Inches(0.5), color=ACCENT):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    bar.fill.solid()
    bar.fill.fore_color.rgb = color
    bar.line.fill.background()
    bar.shadow.inherit = False
    return bar


def add_header(slide, title, subtitle=None):
    accent_bar(slide, Inches(0.5), Inches(0.45), w=Inches(0.08), h=Inches(0.6))
    add_text(slide, Inches(0.7), Inches(0.4), Inches(11), Inches(0.6),
             title, size=28, color=TEXT_WHITE, bold=True)
    if subtitle:
        add_text(slide, Inches(0.7), Inches(0.95), Inches(11), Inches(0.4),
                 subtitle, size=14, color=ACCENT)
    # Bottom rule
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                  Inches(0.5), Inches(1.45),
                                  Inches(12.3), Emu(15000))
    line.fill.solid()
    line.fill.fore_color.rgb = LINE_BLUE
    line.line.fill.background()
    line.shadow.inherit = False


def add_footer(slide, page_num):
    add_text(slide, Inches(0.5), Inches(7.05), Inches(6), Inches(0.3),
             "CN6035 — Theatre Reservation App", size=10, color=TEXT_GRAY)
    add_text(slide, Inches(11.5), Inches(7.05), Inches(1.3), Inches(0.3),
             f"{page_num} / 12", size=10, color=TEXT_GRAY, align=PP_ALIGN.RIGHT)


# ─────────────────────────────────────────────────────────
# SLIDE 1 — Title
# ─────────────────────────────────────────────────────────
s = prs.slides.add_slide(blank_layout)
set_bg(s)

# Centered logo block
logo_y = Inches(1.8)
add_text(s, Inches(0), logo_y, SLIDE_W, Inches(0.8),
         "🎭", size=64, color=ACCENT, align=PP_ALIGN.CENTER)
add_text(s, Inches(0), Inches(2.7), SLIDE_W, Inches(0.5),
         "CN6035", size=20, color=ACCENT, bold=True, align=PP_ALIGN.CENTER)
add_text(s, Inches(0), Inches(3.2), SLIDE_W, Inches(1.2),
         "Theatre Seat Reservation App",
         size=44, color=TEXT_WHITE, bold=True, align=PP_ALIGN.CENTER)
add_text(s, Inches(0), Inches(4.4), SLIDE_W, Inches(0.5),
         "Mobile & Distributed Systems",
         size=18, color=TEXT_GRAY, align=PP_ALIGN.CENTER)

# Decorative line
line = s.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                          Inches(5.5), Inches(5.0),
                          Inches(2.3), Emu(20000))
line.fill.solid(); line.fill.fore_color.rgb = ACCENT
line.line.fill.background(); line.shadow.inherit = False

# Student info
add_text(s, Inches(0), Inches(5.5), SLIDE_W, Inches(0.4),
         "Φοιτητής", size=12, color=TEXT_GRAY, align=PP_ALIGN.CENTER)
add_text(s, Inches(0), Inches(5.85), SLIDE_W, Inches(0.5),
         "[Ονοματεπώνυμο]",
         size=20, color=TEXT_WHITE, bold=True, align=PP_ALIGN.CENTER)
add_text(s, Inches(0), Inches(6.4), SLIDE_W, Inches(0.4),
         "[ΑΜ Φοιτητή]   ·   Ακαδημαϊκό Έτος 2025-2026",
         size=12, color=TEXT_GRAY, align=PP_ALIGN.CENTER)


# ─────────────────────────────────────────────────────────
# SLIDE 2 — Στόχος εφαρμογής
# ─────────────────────────────────────────────────────────
s = prs.slides.add_slide(blank_layout)
set_bg(s)
add_header(s, "Στόχος Εφαρμογής", "Τι λύνει το app και για ποιους")

# Two-column layout
add_box(s, Inches(0.5), Inches(1.8), Inches(6), Inches(5.0))
add_text(s, Inches(0.8), Inches(2.0), Inches(5.5), Inches(0.5),
         "🎯 Πρόβλημα", size=18, color=ACCENT, bold=True)
add_text(s, Inches(0.8), Inches(2.6), Inches(5.5), Inches(4.0),
         "Η αγορά εισιτηρίων για θεατρικές παραστάσεις απαιτεί συχνά:\n"
         "• επίσκεψη στο ταμείο του θεάτρου\n"
         "• πολλαπλά websites με διαφορετικά UI\n"
         "• αδυναμία προβολής διαθέσιμων θέσεων σε πραγματικό χρόνο\n"
         "• κίνδυνο διπλο-κράτησης (race conditions)\n\n"
         "Στόχος: ενιαία εμπειρία από κινητό, παντού, με εγγυημένη ακεραιότητα.",
         size=13, color=TEXT_WHITE)

add_box(s, Inches(6.83), Inches(1.8), Inches(6), Inches(5.0))
add_text(s, Inches(7.13), Inches(2.0), Inches(5.5), Inches(0.5),
         "✅ Λύση", size=18, color=ACCENT, bold=True)
add_text(s, Inches(7.13), Inches(2.6), Inches(5.5), Inches(4.0),
         "Mobile εφαρμογή Theatre Seat Reservation με:\n"
         "• Εγγραφή / Σύνδεση χρηστών (JWT)\n"
         "• Περιήγηση θεάτρων & παραστάσεων\n"
         "• Επιλογή πολλαπλών θέσεων με διαφορετικές κατηγορίες\n"
         "• Ατομικές κρατήσεις με μοναδικό κωδικό αναφοράς\n"
         "• Διαχείριση προφίλ, ιστορικού, αλλαγής/ακύρωσης\n"
         "• Cloud backend (Railway) — προσβάσιμο από παντού",
         size=13, color=TEXT_WHITE)

add_footer(s, 2)


# ─────────────────────────────────────────────────────────
# SLIDE 3 — Tech Stack
# ─────────────────────────────────────────────────────────
s = prs.slides.add_slide(blank_layout)
set_bg(s)
add_header(s, "Tech Stack", "Τεχνολογίες & Εργαλεία")

stack_items = [
    ("📱", "Frontend", "React Native\nExpo SDK 54", "Cross-platform mobile UI"),
    ("⚙️", "Backend", "Node.js\nExpress 4.x", "REST API server"),
    ("🗄", "Database", "MySQL 8.0\n(MariaDB compatible)", "Persistent storage"),
    ("🔐", "Auth", "JWT\n+ bcrypt", "Stateless authentication"),
    ("☁️", "Hosting", "Railway\n+ GitHub CI/CD", "Auto-deploy on push"),
    ("🛠", "Tools", "Postman, Git\nExpo Go, npm", "Development & testing"),
]

cols = 3
card_w = Inches(3.95)
card_h = Inches(2.4)
gap_x = Inches(0.15)
gap_y = Inches(0.25)
start_x = Inches(0.5)
start_y = Inches(1.85)

for idx, (icon, title, tech, desc) in enumerate(stack_items):
    row = idx // cols
    col = idx % cols
    x = start_x + col * (card_w + gap_x)
    y = start_y + row * (card_h + gap_y)
    add_box(s, x, y, card_w, card_h)
    accent_bar(s, x, y, w=card_w, h=Inches(0.05))
    add_text(s, x + Inches(0.3), y + Inches(0.2), Inches(1), Inches(0.6),
             icon, size=32, color=ACCENT)
    add_text(s, x + Inches(1.2), y + Inches(0.3), card_w - Inches(1.4), Inches(0.4),
             title, size=14, color=TEXT_GRAY, bold=True)
    add_text(s, x + Inches(1.2), y + Inches(0.7), card_w - Inches(1.4), Inches(0.7),
             tech, size=14, color=TEXT_WHITE, bold=True)
    add_text(s, x + Inches(0.3), y + Inches(1.6), card_w - Inches(0.6), Inches(0.7),
             desc, size=11, color=TEXT_GRAY)

add_footer(s, 3)


# ─────────────────────────────────────────────────────────
# SLIDE 4 — Αρχιτεκτονική (3-tier)
# ─────────────────────────────────────────────────────────
s = prs.slides.add_slide(blank_layout)
set_bg(s)
add_header(s, "Αρχιτεκτονική Συστήματος", "3-tier κατανεμημένο σύστημα")

# 3 tiers as horizontal blocks
tier_y = Inches(2.3)
tier_h = Inches(2.5)

# Tier 1 — Mobile Client
t1_x = Inches(0.6)
t1_w = Inches(3.5)
add_box(s, t1_x, tier_y, t1_w, tier_h)
accent_bar(s, t1_x, tier_y, w=t1_w, h=Inches(0.05))
add_text(s, t1_x, tier_y + Inches(0.2), t1_w, Inches(0.5),
         "📱  CLIENT", size=14, color=ACCENT, bold=True, align=PP_ALIGN.CENTER)
add_text(s, t1_x, tier_y + Inches(0.7), t1_w, Inches(0.5),
         "iPhone / Android",
         size=18, color=TEXT_WHITE, bold=True, align=PP_ALIGN.CENTER)
add_text(s, t1_x + Inches(0.3), tier_y + Inches(1.3), t1_w - Inches(0.6), Inches(1.1),
         "• React Native (Expo)\n"
         "• expo-secure-store για JWT\n"
         "• Axios HTTP client\n"
         "• React Navigation",
         size=11, color=TEXT_WHITE)

# Tier 2 — REST API
t2_x = Inches(4.9)
t2_w = Inches(3.5)
add_box(s, t2_x, tier_y, t2_w, tier_h)
accent_bar(s, t2_x, tier_y, w=t2_w, h=Inches(0.05))
add_text(s, t2_x, tier_y + Inches(0.2), t2_w, Inches(0.5),
         "⚙️  API SERVER", size=14, color=ACCENT, bold=True, align=PP_ALIGN.CENTER)
add_text(s, t2_x, tier_y + Inches(0.7), t2_w, Inches(0.5),
         "Node.js / Express",
         size=18, color=TEXT_WHITE, bold=True, align=PP_ALIGN.CENTER)
add_text(s, t2_x + Inches(0.3), tier_y + Inches(1.3), t2_w - Inches(0.6), Inches(1.1),
         "• REST endpoints (JSON)\n"
         "• JWT middleware\n"
         "• mysql2 connection pool\n"
         "• bcrypt για passwords",
         size=11, color=TEXT_WHITE)

# Tier 3 — Database
t3_x = Inches(9.2)
t3_w = Inches(3.5)
add_box(s, t3_x, tier_y, t3_w, tier_h)
accent_bar(s, t3_x, tier_y, w=t3_w, h=Inches(0.05))
add_text(s, t3_x, tier_y + Inches(0.2), t3_w, Inches(0.5),
         "🗄  DATABASE", size=14, color=ACCENT, bold=True, align=PP_ALIGN.CENTER)
add_text(s, t3_x, tier_y + Inches(0.7), t3_w, Inches(0.5),
         "MySQL (Railway)",
         size=18, color=TEXT_WHITE, bold=True, align=PP_ALIGN.CENTER)
add_text(s, t3_x + Inches(0.3), tier_y + Inches(1.3), t3_w - Inches(0.6), Inches(1.1),
         "• 6 tables με FK / CASCADE\n"
         "• Transactions (FOR UPDATE)\n"
         "• ACID guarantees\n"
         "• Cloud-hosted",
         size=11, color=TEXT_WHITE)

# Arrows (using simple arrows between tiers)
def arrow(slide, x1, y, x2, label):
    line = slide.shapes.add_connector(1, x1, y, x2, y)
    line.line.color.rgb = ACCENT
    line.line.width = Pt(2.5)
    # Arrow head
    head = slide.shapes.add_shape(MSO_SHAPE.RIGHT_TRIANGLE,
                                  x2 - Inches(0.15), y - Inches(0.07),
                                  Inches(0.18), Inches(0.14))
    head.fill.solid()
    head.fill.fore_color.rgb = ACCENT
    head.line.fill.background()
    head.shadow.inherit = False
    # Label
    add_text(slide, x1, y - Inches(0.45),
             x2 - x1, Inches(0.3),
             label, size=10, color=TEXT_GRAY, align=PP_ALIGN.CENTER, bold=True)

arrow(s, Inches(4.1), tier_y + Inches(1.25), Inches(4.9), "HTTPS / JSON")
arrow(s, Inches(8.4), tier_y + Inches(1.25), Inches(9.2), "TCP / SQL")

# Bottom note
add_box(s, Inches(0.6), Inches(5.3), Inches(12.1), Inches(1.3))
add_text(s, Inches(0.9), Inches(5.5), Inches(11.5), Inches(0.4),
         "🌐  Δικτυακή Επικοινωνία", size=14, color=ACCENT, bold=True)
add_text(s, Inches(0.9), Inches(5.95), Inches(11.5), Inches(0.7),
         "Client ↔ API: HTTPS — Public URL: cn6035-production.up.railway.app  ·  "
         "API ↔ DB: εσωτερικό Railway network (κρυπτογραφημένο, χωρίς exposure)",
         size=11, color=TEXT_WHITE)

add_footer(s, 4)


# ─────────────────────────────────────────────────────────
# SLIDE 5 — Database Schema
# ─────────────────────────────────────────────────────────
s = prs.slides.add_slide(blank_layout)
set_bg(s)
add_header(s, "Σχεδιασμός Βάσης Δεδομένων", "6 πίνακες με σχέσεις FK / CASCADE")

# Tables in 2 rows of 3
tables = [
    ("users", "user_id PK\nname\nemail UNIQUE\npassword (bcrypt)\ncreated_at", "👤"),
    ("theatres", "theatre_id PK\nname\nlocation\ndescription", "🏛"),
    ("shows", "show_id PK\ntheatre_id FK→\ntitle\ndescription\nduration\nage_rating", "🎭"),
    ("showtimes", "showtime_id PK\nshow_id FK→\ndate, time, room\ntotal_seats\navailable_seats\nprice", "🕐"),
    ("seats", "seat_id PK\nshowtime_id FK→\nseat_number\ncategory\nis_available", "💺"),
    ("reservations", "reservation_id PK\nuser_id FK→\nshowtime_id FK→\nseat_id FK→\nreservation_reference\nstatus", "🎫"),
]

card_w = Inches(4.0)
card_h = Inches(2.4)
gap = Inches(0.15)
start_x = Inches(0.5)
start_y = Inches(1.85)

for idx, (name, fields, icon) in enumerate(tables):
    row = idx // 3
    col = idx % 3
    x = start_x + col * (card_w + gap)
    y = start_y + row * (card_h + gap)
    add_box(s, x, y, card_w, card_h)
    accent_bar(s, x, y, w=Inches(0.06), h=card_h, color=ACCENT)
    add_text(s, x + Inches(0.2), y + Inches(0.15), Inches(0.6), Inches(0.5),
             icon, size=20, color=ACCENT)
    add_text(s, x + Inches(0.85), y + Inches(0.2), card_w - Inches(1), Inches(0.5),
             name, size=16, color=TEXT_WHITE, bold=True, font='Consolas')
    add_text(s, x + Inches(0.2), y + Inches(0.75), card_w - Inches(0.4), Inches(1.6),
             fields, size=10, color=TEXT_GRAY, font='Consolas')

# Note about CASCADE
add_text(s, Inches(0.5), Inches(6.85), Inches(12.3), Inches(0.3),
         "Όλες οι FK relations έχουν ON DELETE CASCADE  ·  "
         "Index στο reservation_reference για γρήγορο grouping",
         size=11, color=ACCENT, align=PP_ALIGN.CENTER)

add_footer(s, 5)


# ─────────────────────────────────────────────────────────
# SLIDE 6 — Authentication Flow (JWT)
# ─────────────────────────────────────────────────────────
s = prs.slides.add_slide(blank_layout)
set_bg(s)
add_header(s, "Authentication Flow", "JWT-based, stateless")

# Two-column layout
# Left: register/login flow
add_box(s, Inches(0.5), Inches(1.85), Inches(6.1), Inches(4.95))
add_text(s, Inches(0.8), Inches(2.05), Inches(5.5), Inches(0.5),
         "🔐 Register / Login", size=16, color=ACCENT, bold=True)

steps = [
    ("1.", "User submits", "email + password"),
    ("2.", "Backend hashes password", "bcrypt (cost 10)"),
    ("3.", "INSERT user OR validate", "users table"),
    ("4.", "Sign JWT token", "{userId, email}, 7-day expiry"),
    ("5.", "Return token", "{ token, user }"),
    ("6.", "Client stores token", "expo-secure-store (encrypted)"),
]
for i, (n, t, sub) in enumerate(steps):
    yy = Inches(2.6 + i*0.55)
    add_text(s, Inches(0.8), yy, Inches(0.4), Inches(0.5),
             n, size=14, color=ACCENT, bold=True, font='Consolas')
    add_text(s, Inches(1.2), yy, Inches(3.2), Inches(0.5),
             t, size=12, color=TEXT_WHITE, bold=True)
    add_text(s, Inches(4.4), yy, Inches(2.0), Inches(0.5),
             sub, size=11, color=TEXT_GRAY, font='Consolas')

# Right: protected request flow
add_box(s, Inches(6.83), Inches(1.85), Inches(6.0), Inches(4.95))
add_text(s, Inches(7.1), Inches(2.05), Inches(5.5), Inches(0.5),
         "🛡 Protected API Request", size=16, color=ACCENT, bold=True)

req_steps = [
    ("1.", "Client attaches header", "Authorization: Bearer <jwt>"),
    ("2.", "authMiddleware verifies", "jsonwebtoken.verify(secret)"),
    ("3.", "Extract userId", "req.user.userId"),
    ("4.", "Endpoint handler runs", "reservations / profile / etc."),
    ("5.", "Token expires? → 401", "Auto re-login on client"),
    ("6.", "JWT_SECRET in env vars", "Never in code/git"),
]
for i, (n, t, sub) in enumerate(req_steps):
    yy = Inches(2.6 + i*0.55)
    add_text(s, Inches(7.1), yy, Inches(0.4), Inches(0.5),
             n, size=14, color=ACCENT, bold=True, font='Consolas')
    add_text(s, Inches(7.5), yy, Inches(3.2), Inches(0.5),
             t, size=12, color=TEXT_WHITE, bold=True)
    add_text(s, Inches(10.7), yy, Inches(2.1), Inches(0.5),
             sub, size=10, color=TEXT_GRAY, font='Consolas')

add_footer(s, 6)


# ─────────────────────────────────────────────────────────
# SLIDE 7 — REST API Endpoints
# ─────────────────────────────────────────────────────────
s = prs.slides.add_slide(blank_layout)
set_bg(s)
add_header(s, "REST API Endpoints", "12 endpoints συνολικά — όλα με JSON responses")

endpoints = [
    ("POST",  "/api/register",                "—",   "Δημιουργία νέου χρήστη"),
    ("POST",  "/api/login",                   "—",   "Authenticate, επιστρέφει JWT"),
    ("GET",   "/api/theatres",                "—",   "Λίστα θεάτρων (?search=...)"),
    ("GET",   "/api/theatres/:id",            "—",   "Λεπτομέρειες θεάτρου"),
    ("GET",   "/api/shows",                   "—",   "Παραστάσεις (theatreId, title, date)"),
    ("GET",   "/api/shows/:id/showtimes",     "—",   "Showtimes συγκεκριμένης παράστασης"),
    ("GET",   "/api/showtimes/:id/seats",     "JWT", "Διαθεσιμότητα θέσεων"),
    ("POST",  "/api/reservations",            "JWT", "Multi-seat κράτηση (atomic)"),
    ("PUT",   "/api/reservations/:id",        "JWT", "Αλλαγή θέσης"),
    ("DELETE","/api/reservations/:id",        "JWT", "Ακύρωση κράτησης"),
    ("GET",   "/api/user/reservations",       "JWT", "Κρατήσεις χρήστη"),
    ("GET",   "/api/reservations/my",         "JWT", "Alias του παραπάνω"),
]

# Header row
hdr_y = Inches(1.85)
add_text(s, Inches(0.6), hdr_y, Inches(1), Inches(0.4),
         "METHOD", size=11, color=ACCENT, bold=True, font='Consolas')
add_text(s, Inches(1.7), hdr_y, Inches(4.2), Inches(0.4),
         "ENDPOINT", size=11, color=ACCENT, bold=True, font='Consolas')
add_text(s, Inches(6.0), hdr_y, Inches(0.8), Inches(0.4),
         "AUTH", size=11, color=ACCENT, bold=True, font='Consolas')
add_text(s, Inches(7.0), hdr_y, Inches(6), Inches(0.4),
         "DESCRIPTION", size=11, color=ACCENT, bold=True, font='Consolas')

# Separator
sep = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(2.25),
                        Inches(12.3), Emu(8000))
sep.fill.solid(); sep.fill.fore_color.rgb = LINE_BLUE
sep.line.fill.background(); sep.shadow.inherit = False

# Color map for methods
method_color = {
    "GET": ACCENT_GREEN,
    "POST": ACCENT,
    "PUT": RGBColor(0x60, 0xa5, 0xfa),
    "DELETE": ACCENT_RED,
}

for i, (m, ep, auth, desc) in enumerate(endpoints):
    yy = Inches(2.4 + i*0.38)
    # Row background (alternate)
    if i % 2 == 1:
        bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                Inches(0.5), yy - Inches(0.05),
                                Inches(12.3), Inches(0.38))
        bg.fill.solid(); bg.fill.fore_color.rgb = BG_CARD
        bg.line.fill.background(); bg.shadow.inherit = False

    add_text(s, Inches(0.6), yy, Inches(1), Inches(0.4),
             m, size=11, color=method_color.get(m, TEXT_WHITE), bold=True, font='Consolas')
    add_text(s, Inches(1.7), yy, Inches(4.2), Inches(0.4),
             ep, size=11, color=TEXT_WHITE, font='Consolas')
    add_text(s, Inches(6.0), yy, Inches(0.8), Inches(0.4),
             auth, size=10, color=ACCENT if auth == "JWT" else TEXT_GRAY, bold=(auth=="JWT"), font='Consolas')
    add_text(s, Inches(7.0), yy, Inches(6), Inches(0.4),
             desc, size=11, color=TEXT_GRAY)

add_footer(s, 7)


# ─────────────────────────────────────────────────────────
# SLIDE 8 — Multi-Seat Booking & Concurrency
# ─────────────────────────────────────────────────────────
s = prs.slides.add_slide(blank_layout)
set_bg(s)
add_header(s, "Multi-Seat Booking & Concurrency",
           "Atomic transactions, FOR UPDATE locks, booking references")

# Code-like flow on left
add_box(s, Inches(0.5), Inches(1.85), Inches(7.0), Inches(4.95), fill=BG_CARD)
add_text(s, Inches(0.75), Inches(2.05), Inches(6.5), Inches(0.4),
         "📥  POST /api/reservations  { showtime_id, seat_ids: [11, 12, 27] }",
         size=12, color=ACCENT, bold=True, font='Consolas')

steps = [
    "1. Validation: positive integers, no duplicates, ≤10",
    "2. BEGIN TRANSACTION",
    "3. SELECT seats WHERE seat_id IN (...) FOR UPDATE",
    "   → row-level locks αποτρέπουν παράλληλη κράτηση",
    "4. Έλεγχος: όλες οι θέσεις διαθέσιμες;",
    "   → αν όχι, ROLLBACK + 409 Conflict",
    "5. generateReference() → BK-20260424-A7X3F2",
    "   crypto.randomBytes(3) — 16.7M unique codes/day",
    "6. INSERT N rows με ίδιο reservation_reference",
    "7. UPDATE seats SET is_available = FALSE",
    "8. UPDATE showtimes SET available_seats = ... - N",
    "9. COMMIT  ✅",
    "10. Response: { reference, count, totalPrice }",
]
for i, t in enumerate(steps):
    yy = Inches(2.55 + i*0.32)
    color = TEXT_WHITE if not t.startswith("   ") else TEXT_GRAY
    weight = True if not t.startswith("   ") else False
    add_text(s, Inches(0.85), yy, Inches(6.6), Inches(0.35),
             t, size=11, color=color, bold=weight, font='Consolas')

# Right side — guarantees
add_box(s, Inches(7.83), Inches(1.85), Inches(5.0), Inches(4.95))
add_text(s, Inches(8.1), Inches(2.05), Inches(4.5), Inches(0.4),
         "🛡 Εγγυήσεις",
         size=16, color=ACCENT, bold=True)

guarantees = [
    ("Atomicity", "Είτε όλες οι θέσεις κρατούνται, είτε καμία"),
    ("No Double-Booking", "FOR UPDATE locks σειριοποιούν concurrent αιτήματα"),
    ("Consistency", "available_seats counter πάντα συγχρονισμένος"),
    ("Validated Input", "Δεν περνά κακόβουλο payload"),
    ("Pricing on Server", "Τιμές υπολογίζονται από category, όχι client"),
    ("Tracking", "Booking reference για αναφορά/υποστήριξη"),
]
for i, (t, d) in enumerate(guarantees):
    yy = Inches(2.6 + i*0.7)
    add_text(s, Inches(8.1), yy, Inches(4.5), Inches(0.35),
             "✓ " + t, size=12, color=ACCENT_GREEN, bold=True)
    add_text(s, Inches(8.3), yy + Inches(0.32), Inches(4.5), Inches(0.4),
             d, size=10, color=TEXT_GRAY)

add_footer(s, 8)


# ─────────────────────────────────────────────────────────
# SLIDE 9 — Screenshots: Login + Theatres
# ─────────────────────────────────────────────────────────
s = prs.slides.add_slide(blank_layout)
set_bg(s)
add_header(s, "Δείγματα Χρήσης (1/3)", "Login & Home Screen")

# Two screenshot placeholders
def screenshot_placeholder(slide, x, y, w, h, label, desc):
    add_box(slide, x, y, w, h, fill=BG_CARD, line=ACCENT, line_w=2)
    add_text(slide, x, y + h/2 - Inches(0.4), w, Inches(0.4),
             "📸", size=36, color=ACCENT, align=PP_ALIGN.CENTER)
    add_text(slide, x, y + h/2 + Inches(0.1), w, Inches(0.4),
             label, size=14, color=TEXT_GRAY, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, x, y + h/2 + Inches(0.5), w, Inches(0.4),
             "[Insert screenshot here]", size=10, color=TEXT_GRAY, align=PP_ALIGN.CENTER)

# 2 phone-shaped placeholders side by side
phone_w = Inches(2.6)
phone_h = Inches(5.2)

# Left
screenshot_placeholder(s, Inches(1.5), Inches(1.85), phone_w, phone_h,
                       "Login Screen", "")
add_text(s, Inches(0.5), Inches(7.15), Inches(4.5), Inches(0.3),
         "🔐 Είσοδος / Εγγραφή με JWT",
         size=12, color=ACCENT, bold=True, align=PP_ALIGN.CENTER)

# Center description card
add_box(s, Inches(4.85), Inches(1.85), Inches(3.6), phone_h, fill=BG_CARD)
add_text(s, Inches(5.1), Inches(2.1), Inches(3.1), Inches(0.5),
         "Λειτουργίες",
         size=16, color=ACCENT, bold=True)
features = [
    "• Email + password validation",
    "• Bcrypt password hashing",
    "• JWT 7-ημερών",
    "• expo-secure-store (encrypted)",
    "",
    "Home Screen",
    "• Hero banner με χαιρετισμό",
    "• Stats: 6 θέατρα / πόλεις / 14 παρ.",
    "• Filter chips (Τραγωδία, Κωμωδία...)",
    "• Search box (theatre / location)",
    "• Pull-to-refresh",
    "• Theatre cards με 'next show' badges",
]
for i, f in enumerate(features):
    yy = Inches(2.65 + i*0.32)
    color = ACCENT if f and not f.startswith('•') else (TEXT_WHITE if f else TEXT_WHITE)
    bold = True if f and not f.startswith('•') else False
    if not f:
        continue
    add_text(s, Inches(5.1), yy, Inches(3.3), Inches(0.35),
             f, size=11, color=color, bold=bold)

# Right
screenshot_placeholder(s, Inches(8.85), Inches(1.85), phone_w, phone_h,
                       "Theatre List Screen", "")
add_text(s, Inches(7.85), Inches(7.15), Inches(4.5), Inches(0.3),
         "🏛 Λίστα Θεάτρων + Search + Filters",
         size=12, color=ACCENT, bold=True, align=PP_ALIGN.CENTER)

add_footer(s, 9)


# ─────────────────────────────────────────────────────────
# SLIDE 10 — Screenshots: Reservation Flow
# ─────────────────────────────────────────────────────────
s = prs.slides.add_slide(blank_layout)
set_bg(s)
add_header(s, "Δείγματα Χρήσης (2/3)", "Show Detail & Multi-Seat Booking")

# 3 phone placeholders
phone_w = Inches(2.4)
phone_h = Inches(4.8)

screenshot_placeholder(s, Inches(0.5), Inches(1.85), phone_w, phone_h,
                       "Show Detail", "")
add_text(s, Inches(0.5), Inches(6.75), phone_w, Inches(0.3),
         "Showtimes + Director", size=11, color=ACCENT, bold=True, align=PP_ALIGN.CENTER)

screenshot_placeholder(s, Inches(3.15), Inches(1.85), phone_w, phone_h,
                       "Seat Selection", "")
add_text(s, Inches(3.15), Inches(6.75), phone_w, Inches(0.3),
         "Multi-seat + 3 sections", size=11, color=ACCENT, bold=True, align=PP_ALIGN.CENTER)

screenshot_placeholder(s, Inches(5.8), Inches(1.85), phone_w, phone_h,
                       "Price Breakdown", "")
add_text(s, Inches(5.8), Inches(6.75), phone_w, Inches(0.3),
         "Live total + categories", size=11, color=ACCENT, bold=True, align=PP_ALIGN.CENTER)

# Description card on right
add_box(s, Inches(8.5), Inches(1.85), Inches(4.3), phone_h)
add_text(s, Inches(8.75), Inches(2.05), Inches(3.8), Inches(0.5),
         "🎟  Booking Flow",
         size=16, color=ACCENT, bold=True)
flow_text = [
    "1. Επιλογή showtime",
    "2. Grid 144 θέσεων (12 σειρές × 12)",
    "   • A-B → VIP",
    "   • C-H → Standard",
    "   • I-L → Economy",
    "3. Tap to select / deselect",
    "4. Live breakdown:",
    "   2× VIP (€37.50) = €75.00",
    "   1× Standard (€25) = €25.00",
    "   Σύνολο: €100",
    "5. Confirm → POST /reservations",
    "6. Success με reference code",
    "7. Auto navigate σε Profile",
]
for i, t in enumerate(flow_text):
    yy = Inches(2.6 + i*0.32)
    color = TEXT_WHITE if not t.startswith("   ") else TEXT_GRAY
    add_text(s, Inches(8.75), yy, Inches(3.8), Inches(0.35),
             t, size=10, color=color, font='Consolas')

add_footer(s, 10)


# ─────────────────────────────────────────────────────────
# SLIDE 11 — Screenshots: Profile + Reservation Management
# ─────────────────────────────────────────────────────────
s = prs.slides.add_slide(blank_layout)
set_bg(s)
add_header(s, "Δείγματα Χρήσης (3/3)", "Profile & Reservation Management")

phone_w = Inches(2.6)
phone_h = Inches(5.2)

# Left — Profile main
screenshot_placeholder(s, Inches(0.7), Inches(1.85), phone_w, phone_h,
                       "Profile Screen", "")
add_text(s, Inches(0.7), Inches(7.15), phone_w, Inches(0.3),
         "Tabs + Stats + Bookings",
         size=11, color=ACCENT, bold=True, align=PP_ALIGN.CENTER)

# Middle description
add_box(s, Inches(3.7), Inches(1.85), Inches(5.5), phone_h)
add_text(s, Inches(3.95), Inches(2.05), Inches(5.0), Inches(0.5),
         "👤  Λειτουργίες Profile",
         size=16, color=ACCENT, bold=True)

prof_text = [
    ("Επερχόμενες / Ιστορικό",  "tabs με grouped bookings"),
    ("Booking Reference",       "🎫 BK-20260424-A7X3F2"),
    ("Per-seat actions",        "Αλλαγή θέσης / Ακύρωση"),
    ("Mixed status",            "ακυρωμένες θέσεις με strikethrough"),
    ("Booking total",           "δυναμικός υπολογισμός μετά από ακύρωση"),
    ("Stats row",               "συνολικός αριθμός κρατήσεων"),
    ("Pull-to-refresh",         "ανανέωση από server"),
    ("Logout",                  "διαγραφή JWT από secure store"),
]
for i, (t, d) in enumerate(prof_text):
    yy = Inches(2.65 + i*0.55)
    add_text(s, Inches(3.95), yy, Inches(2.3), Inches(0.4),
             "▸ " + t, size=12, color=TEXT_WHITE, bold=True)
    add_text(s, Inches(6.3), yy, Inches(2.85), Inches(0.4),
             d, size=11, color=TEXT_GRAY)

# Right — Reservation card detail (placeholder)
screenshot_placeholder(s, Inches(9.6), Inches(1.85), phone_w, phone_h,
                       "Booking Card Detail", "")
add_text(s, Inches(9.6), Inches(7.15), phone_w, Inches(0.3),
         "Reference + Seats list",
         size=11, color=ACCENT, bold=True, align=PP_ALIGN.CENTER)

add_footer(s, 11)


# ─────────────────────────────────────────────────────────
# SLIDE 12 — Συμπεράσματα & Links
# ─────────────────────────────────────────────────────────
s = prs.slides.add_slide(blank_layout)
set_bg(s)
add_header(s, "Συμπεράσματα", "Επίτευξη στόχων & μελλοντικές επεκτάσεις")

# Left — Επιτεύγματα
add_box(s, Inches(0.5), Inches(1.85), Inches(6.1), Inches(3.4))
add_text(s, Inches(0.8), Inches(2.05), Inches(5.5), Inches(0.5),
         "✅ Επιτεύχθηκαν",
         size=16, color=ACCENT_GREEN, bold=True)
done = [
    "Πλήρες 3-tier mobile app (Client + API + DB)",
    "JWT authentication με bcrypt + secure storage",
    "Multi-seat atomic booking με category pricing",
    "Concurrency safety (FOR UPDATE locks)",
    "Booking reference codes & profile grouping",
    "Cloud deployment (Railway) — runs from anywhere",
]
for i, t in enumerate(done):
    yy = Inches(2.6 + i*0.4)
    add_text(s, Inches(0.85), yy, Inches(0.3), Inches(0.4),
             "✓", size=14, color=ACCENT_GREEN, bold=True)
    add_text(s, Inches(1.15), yy, Inches(5.3), Inches(0.4),
             t, size=12, color=TEXT_WHITE)

# Right — Future work
add_box(s, Inches(6.83), Inches(1.85), Inches(6.0), Inches(3.4))
add_text(s, Inches(7.13), Inches(2.05), Inches(5.5), Inches(0.5),
         "🚀 Μελλοντικές επεκτάσεις",
         size=16, color=ACCENT, bold=True)
future = [
    "OIDC / Keycloak (αντί JWT) για enterprise SSO",
    "Real payment gateway (Stripe / Viva)",
    "QR code tickets για είσοδο στο θέατρο",
    "Push notifications (υπενθυμίσεις)",
    "Seat hold timer (5-min reservation window)",
    "Standalone build (EAS) για App Store / Play Store",
]
for i, t in enumerate(future):
    yy = Inches(2.6 + i*0.4)
    add_text(s, Inches(7.18), yy, Inches(0.3), Inches(0.4),
             "▸", size=14, color=ACCENT, bold=True)
    add_text(s, Inches(7.5), yy, Inches(5.2), Inches(0.4),
             t, size=12, color=TEXT_WHITE)

# Bottom — Links
add_box(s, Inches(0.5), Inches(5.4), Inches(12.3), Inches(1.4))
accent_bar(s, Inches(0.5), Inches(5.4), w=Inches(0.08), h=Inches(1.4))
add_text(s, Inches(0.8), Inches(5.55), Inches(11.5), Inches(0.4),
         "🔗  Πηγαίος Κώδικας & Live Demo",
         size=14, color=ACCENT, bold=True)
add_text(s, Inches(0.8), Inches(5.95), Inches(2), Inches(0.4),
         "GitHub:", size=12, color=TEXT_GRAY, bold=True)
add_text(s, Inches(2.3), Inches(5.95), Inches(7), Inches(0.4),
         "github.com/Yungmadara/CN6035", size=12, color=TEXT_WHITE, font='Consolas')
add_text(s, Inches(0.8), Inches(6.4), Inches(2), Inches(0.4),
         "Live API:", size=12, color=TEXT_GRAY, bold=True)
add_text(s, Inches(2.3), Inches(6.4), Inches(8), Inches(0.4),
         "cn6035-production.up.railway.app", size=12, color=TEXT_WHITE, font='Consolas')

# Thank you
add_text(s, Inches(0.5), Inches(7.0), SLIDE_W - Inches(1), Inches(0.4),
         "🙏  Ευχαριστώ για την προσοχή σας  ·  Ερωτήσεις;",
         size=14, color=ACCENT, bold=True, align=PP_ALIGN.CENTER)


# ─────────────────────────────────────────────────────────
# SAVE
# ─────────────────────────────────────────────────────────
out_dir = os.path.dirname(__file__)
out_path = os.path.join(out_dir, '..', 'CN6035_Presentation.pptx')
out_path = os.path.normpath(out_path)
prs.save(out_path)
print(f"Generated: {out_path}")
print(f"Slides: {len(prs.slides)}")
